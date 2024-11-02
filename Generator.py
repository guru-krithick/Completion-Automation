from docx import Document
from docx2pdf import convert
import os

# Function to replace text in the template
def replace_text_in_docx(input_docx_path, output_docx_path, replacements):
    doc = Document(input_docx_path)

    # Replace <NAME> and <DOMAIN> in the document
    for paragraph in doc.paragraphs:
        for old, new in replacements.items():
            if old in paragraph.text:
                for run in paragraph.runs:
                    run.text = run.text.replace(old, new)

    # Save the modified document
    doc.save(output_docx_path)

# Function to convert DOCX to PDF
def convert_docx_to_pdf(input_docx_path, output_pdf_path):
    convert(input_docx_path, output_pdf_path)

def generate_certificate(row, template_file, output_dir='.'):
    # Extract the Name and Domain from the row
    row_name = row['Name'].strip()  # Keep the original case of the name, no uppercasing
    row_domain = row['Domain']

    # Create replacements for the placeholders
    replacements = {
        '<NAME>': row_name,
        '<DOMAIN>': row_domain,
    }

    # Construct the output file names directly without adding './' prefix
    output_docx = os.path.join(output_dir, f'{row_name}_Certificate.docx')
    output_pdf = os.path.join(output_dir, f'DNYX-Completion-{row_name}.pdf')

    # Ensure the directory exists
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    # Replace placeholders and generate DOCX
    try:
        replace_text_in_docx(template_file, output_docx, replacements)
    except Exception as e:
        raise e

    # Convert DOCX to PDF
    try:
        convert_docx_to_pdf(output_docx, output_pdf)
    except Exception as e:
        raise e

    # Clean up DOCX file after conversion
    try:
        os.remove(output_docx)
    except Exception as e:
        raise e

    return output_pdf  # Return the correct PDF path 
